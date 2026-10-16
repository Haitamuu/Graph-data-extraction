"""
Genere le PowerPoint de presentation du projet.
Usage : <python avec python-pptx> presentation/build_pptx.py
Les images de resultats sont lues dans ../data (produites par main.py).
"""
import os

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
DATA = os.path.join(os.path.dirname(HERE), "data")
OUT = os.path.join(HERE, "Extraction_automatique_de_courbes_v3.pptx")

NAVY = RGBColor(0x14, 0x23, 0x3C)
TEAL = RGBColor(0x1F, 0x9E, 0x9A)
ORANGE = RGBColor(0xE8, 0x8A, 0x1A)
GREY = RGBColor(0x55, 0x5E, 0x6B)
LIGHT = RGBColor(0xF2, 0xF5, 0xF8)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
RED = RGBColor(0xC0, 0x39, 0x2B)

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
BLANK = prs.slide_layouts[6]


# ------------------------------------------------------------------ helpers
def text_box(slide, x, y, w, h, text, size=18, color=NAVY, bold=False, align=PP_ALIGN.LEFT,
             anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = text
    r.font.size, r.font.bold, r.font.color.rgb = Pt(size), bold, color
    return tb


def bullets(slide, items, x, y, w, h, size=20, color=NAVY):
    """items : liste de str ou (str, niveau)."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    for i, it in enumerate(items):
        txt, lvl = it if isinstance(it, tuple) else (it, 0)
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(9)
        r = p.add_run()
        r.text = ("\u2022 " if lvl == 0 else "\u2013 ") + txt
        r.font.size = Pt(size - 3 * lvl)
        r.font.color.rgb = color if lvl == 0 else GREY
        p.level = lvl
    return tb


def box(slide, x, y, w, h, text, fill=TEAL, size=16, color=WHITE, bold=True, shape=MSO_SHAPE.ROUNDED_RECTANGLE):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    s.line.fill.background()
    tf = s.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    for i, line in enumerate(text.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = line
        r.font.size, r.font.color.rgb = Pt(size if i == 0 else size - 3), color
        r.font.bold = bold if i == 0 else False
    return s


def picture(slide, path, x, y, max_w, max_h):
    """Insere l'image en conservant son ratio dans la zone (x, y, max_w, max_h) ; retourne False si absente."""
    if not os.path.exists(path):
        box(slide, x, y, max_w, max_h, f"image manquante\n{os.path.basename(path)}", fill=RED, size=14)
        return False
    iw, ih = Image.open(path).size
    ratio = min(max_w / iw, max_h / ih)
    w, h = iw * ratio, ih * ratio
    slide.shapes.add_picture(path, Inches(x + (max_w - w) / 2), Inches(y + (max_h - h) / 2), Inches(w), Inches(h))
    return True


def new_slide(title, subtitle=None, notes=""):
    s = prs.slides.add_slide(BLANK)
    band = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Inches(1.15))
    band.fill.solid()
    band.fill.fore_color.rgb = NAVY
    band.line.fill.background()
    text_box(s, 0.5, 0.15, 12.3, 0.7, title, size=30, color=WHITE, bold=True, anchor=MSO_ANCHOR.MIDDLE)
    if subtitle:
        text_box(s, 0.5, 0.72, 12.3, 0.4, subtitle, size=15, color=RGBColor(0xB8, 0xC7, 0xDA))
    accent = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, Inches(1.15), prs.slide_width, Inches(0.06))
    accent.fill.solid()
    accent.fill.fore_color.rgb = TEAL
    accent.line.fill.background()
    text_box(s, 12.3, 7.05, 0.8, 0.3, str(len(prs.slides)), size=11, color=GREY, align=PP_ALIGN.RIGHT)
    s.notes_slide.notes_text_frame.text = notes
    return s


def arrow(slide, x, y, w=0.3, h=0.35, color=ORANGE):
    a = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(x), Inches(y), Inches(w), Inches(h))
    a.fill.solid()
    a.fill.fore_color.rgb = color
    a.line.fill.background()
    return a

def code_box(slide, x, y, w, h, code, size=11, bg=RGBColor(0x28, 0x2A, 0x36), fg=RGBColor(0xF8, 0xF8, 0xF2)):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = bg
    s.line.fill.background()
    tf = s.text_frame
    tf.word_wrap = False
    tf.vertical_anchor = MSO_ANCHOR.TOP
    for i, line in enumerate(code.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        p.space_after = Pt(2)
        r = p.add_run()
        r.text = line
        r.font.name = "Consolas"
        r.font.size = Pt(size)
        r.font.color.rgb = fg
    return s


# ------------------------------------------------------------------ 1. titre
s = prs.slides.add_slide(BLANK)
bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
bg.fill.solid()
bg.fill.fore_color.rgb = NAVY
bg.line.fill.background()
bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(3.55), Inches(1.6), Inches(0.08))
bar.fill.solid()
bar.fill.fore_color.rgb = TEAL
bar.line.fill.background()
text_box(s, 0.8, 1.7, 11.5, 1.8, "Extraction automatique de courbes\nà partir d'images de graphes", size=42,
         color=WHITE, bold=True)
text_box(s, 0.8, 3.8, 11.5, 0.6, "OCR + vision par ordinateur  \u2192  données Excel, sans aucune saisie manuelle",
         size=22, color=RGBColor(0xB8, 0xC7, 0xDA))
text_box(s, 0.8, 6.4, 11.5, 0.5, "Projet technique \u2013 EMSE, 2A", size=16, color=RGBColor(0x8F, 0xA3, 0xBC))
s.notes_slide.notes_text_frame.text = (
    "Présentation du développement de l'outil : de l'idée initiale au pipeline actuel, avec les problèmes "
    "rencontrés et ce qui reste à faire. Durée visée : environ 10 minutes.")

# ------------------------------------------------------------------ 2. contexte
s = new_slide("Contexte et objectif", "Pourquoi extraire des données depuis une image ?", notes=(
    "Beaucoup de données scientifiques ne sont disponibles que sous forme de figures dans des articles. "
    "Les outils existants (WebPlotDigitizer...) demandent de cliquer sur les axes et sur chaque courbe. "
    "Notre exigence : zéro saisie, y compris plusieurs courbes dans la même image."))
bullets(s, [
    "Les résultats publiés sont souvent des figures : la donnée brute n'est pas fournie",
    "Les outils manuels demandent de calibrer les axes et de cliquer sur chaque courbe",
    "Objectif : image en entrée \u2192 fichier Excel en sortie, 100 % automatique",
    ("Plusieurs courbes dans un même graphe, séparées et exportées chacune", 1),
    ("Échelles lues directement dans l'image (OCR), pas saisies par l'utilisateur", 1),
    ("Courbes colorées et courbes noires/grises", 1),
    "Contrainte : une erreur silencieuse est pire qu'un échec signalé",
], 0.7, 1.6, 12, 5, size=24)

# ------------------------------------------------------------------ 3. pipeline
s = new_slide("Architecture du pipeline", "Un seul script : main.py", notes=(
    "Chaque image traverse six étapes. Les deux premières lisent l'image comme un humain (axes, légende), "
    "les suivantes isolent les courbes et convertissent les pixels en valeurs. "
    "La dernière écrit un Excel et une image de contrôle pour vérifier visuellement."))
steps = [("Image", "PNG / JPG"), ("OCR", "texte + chiffres"), ("Calibration", "pixel \u2192 valeur"),
         ("Courbes", "couleur + noir"), ("Légendes", "nom des courbes"), ("Export", "Excel + contrôle")]
x0, bw, gap = 0.5, 1.8, 0.3
for i, (t, sub) in enumerate(steps):
    x = x0 + i * (bw + gap)
    box(s, x, 2.6, bw, 1.4, f"{t}\n{sub}", fill=TEAL if i not in (0, 5) else NAVY, size=19)
    if i < len(steps) - 1:
        arrow(s, x + bw + 0.0, 3.1, gap, 0.35)
text_box(s, 0.5, 4.5, 12.3, 0.5, "Principe : on reproduit ce que fait l'œil humain, étape par étape",
         size=20, color=GREY, bold=True, align=PP_ALIGN.CENTER)
bullets(s, [
    "Bibliothèques : OpenCV (image), EasyOCR (texte), scikit-learn (clustering), pandas (Excel), matplotlib (contrôle)",
    "Dossier images/ en entrée, dossier data/ en sortie (un .xlsx et une image de vérification par graphe)",
], 0.7, 5.2, 12, 1.8, size=18)

# ------------------------------------------------------------------ 3b. history - OCR
s = new_slide("Les débuts : Premières analyses OCR", "Script ocr_ext.py et résultats bruts", notes=(
    "Au début du projet, on a juste lancé EasyOCR sur l'image brute. "
    "Le résultat était un tas de textes sans structure. L'OCR lisait très mal les petits chiffres des axes. "
    "On n'avait aucune correspondance pixel-valeur."))
bullets(s, [
    "Approche initiale : lancer EasyOCR directement sur l'image sans prétraitement (script ocr_ext.py)",
    "Constats de ces premiers tests :",
    ("Les petits chiffres des graduations étaient ignorés ou très mal lus", 1),
    ("Les tirets de graduation étaient confondus avec des signes moins (-)", 1),
    ("Le texte des légendes se mélangeait avec les axes de manière désordonnée", 1),
    ("Aucune notion d'échelle : les textes étaient juste détectés à des coordonnées arbitraires", 1),
    "Conclusion : l'OCR brut ne suffisait pas, il fallait une vraie phase de calibration",
], 0.7, 1.6, 11, 4.5, size=22)

# ------------------------------------------------------------------ 3c. history - Codes initiaux
s = new_slide("Les débuts : Premiers codes d'extraction", "Scripts graphify_main.py et proc.py", notes=(
    "Les premiers essais pour extraire les courbes étaient très naïfs. "
    "On utilisait des seuillages de couleur manuels codés en dur. "
    "Si l'image contenait plusieurs courbes, elles étaient toutes mélangées dans un même nuage de points."))
bullets(s, [
    "Premières tentatives (scripts graphify_main.py, proc.py) :",
    "Méthode naïve : seuillage de couleurs avec des plages codées en dur",
    ("Très fragile face au changement d'éclairage ou d'anti-aliasing", 1),
    "Problème majeur : le mélange des courbes",
    ("Toutes les courbes colorées étaient fusionnées en un seul nuage de points", 1),
    ("L'axe Y restait inversé (l'origine de l'image OpenCV est en haut à gauche)", 1),
], 0.5, 1.6, 7.0, 4.5, size=20)
picture(s, os.path.join(DATA, "bad_result_mixed.png"), 7.6, 1.6, 5.2, 5.0)
text_box(s, 7.6, 6.7, 5.2, 0.5, "Résultat brut : les 3 courbes mélangées, axe inversé", size=14, color=GREY, align=PP_ALIGN.CENTER)

# ------------------------------------------------------------------ 4. OCR + calibration
s = new_slide("Étape 1 \u2013 Lire et calibrer les axes", "OCR des graduations \u2192 relation pixel \u2194 valeur", notes=(
    "L'astuce reprise du dépôt GTE.py : agrandir l'image et la binariser avant l'OCR, sinon les petits chiffres "
    "sont illisibles. Ensuite, les chiffres de l'axe X sont alignés horizontalement et ceux de l'axe Y verticalement. "
    "Un ajustement avec consensus (type RANSAC) écarte les chiffres mal lus. Si moins de 4 graduations sont cohérentes "
    "(ou 3 régulièrement espacées), on n'invente pas d'échelle : on bascule en pixels et on le signale."))
bullets(s, [
    "Agrandissement ×4 (plafonné à 2560 px) + binarisation avant l'OCR",
    "Graduations X = même ordonnée, graduations Y = même abscisse \u2192 regroupement par alignement",
    "Ajustement valeur = a · pixel + b avec rejet des valeurs aberrantes (consensus RANSAC)",
    ("ex. un « 5 » mal lu à la place d'un « 0 » est écarté automatiquement", 1),
    "2e lecture avec un autre prétraitement seulement si la calibration est peu fiable",
    "Calibration refusée si < 4 graduations cohérentes (ou 3 non régulières)",
    ("on bascule en pixels, signalé, plutôt qu'une échelle fausse", 1),
], 0.7, 1.55, 7.6, 5.5, size=21)
box(s, 8.7, 2.0, 4.1, 1.2, "Image\n(chiffres flous)", fill=GREY, size=18)
arrow(s, 10.5, 3.3, 0.5, 0.45, TEAL).rotation = 90
box(s, 8.7, 3.9, 4.1, 1.2, "OCR sur image ×4\n[0, 2, 5, 6, 8, 10]", fill=TEAL, size=18)
arrow(s, 10.5, 5.2, 0.5, 0.45, TEAL).rotation = 90
box(s, 8.7, 5.8, 4.1, 1.2, "Ajustement robuste\nX : 0 \u2192 10", fill=NAVY, size=18)

# ------------------------------------------------------------------ 5. couleurs
s = new_slide("Étape 2 \u2013 Séparer les courbes colorées", "Clustering de teinte (HSV) avec MeanShift", notes=(
    "On travaille sur la teinte et non sur le RGB, parce que la teinte est stable quand la couleur est plus claire "
    "ou plus foncée (anti-aliasing). La teinte est un angle : on la projette sur un cercle (cos, sin) pour que "
    "rouge à 0 et rouge à 179 soient voisins. MeanShift trouve seul le nombre de courbes, contrairement à KMeans."))
bullets(s, [
    "Conversion en HSV, on ne garde que les pixels saturés (S > 50) : le blanc, le noir et le texte sont ignorés",
    "La teinte est un angle \u2192 projection sur un cercle (cos, sin)",
    ("le rouge à 0° et le rouge à 179° deviennent voisins", 1),
    "MeanShift : le nombre de courbes est trouvé automatiquement",
    ("KMeans reste disponible si on veut forcer un nombre", 1),
    "Clusters < 100 pixels écartés (bruit de compression JPEG)",
    "Échantillon fixe de 3000 pixels, graine fixée \u2192 résultats reproductibles",
], 0.7, 1.55, 7.6, 5.5, size=21)
picture(s, os.path.join(DATA, "test_graph_verification.png"), 8.4, 2.0, 4.6, 3.5)
text_box(s, 8.4, 5.6, 4.6, 0.8, "test_graph : courbes bleue et rouge séparées automatiquement", size=14, color=GREY,
         align=PP_ALIGN.CENTER)

# ------------------------------------------------------------------ 6. noir
s = new_slide("Étape 3 \u2013 Extraire la courbe noire", "Difficile : le texte, les axes et la grille sont noirs aussi",
              notes=(
                  "Par soustraction : on part de tous les pixels sombres et on enlève tout ce qui n'est pas la courbe. "
                  "Le texte est effacé grâce aux boîtes de l'OCR, les axes et la grille avec une ouverture morphologique "
                  "(un trait horizontal ou vertical d'au moins 40 pixels), puis les courbes colorées. Ce qui reste est "
                  "la courbe. Piège rencontré : contourArea vaut 0 pour une ligne d'un pixel, il faut arcLength."))
steps = [("Pixels sombres", GREY), ("– texte\n(boîtes OCR)", TEAL), ("– axes / grille\n(morphologie)", TEAL),
         ("– courbes colorées\n(masque dilaté)", TEAL), ("Courbe noire", NAVY)]
x0, bw, gap = 0.5, 2.1, 0.45
for i, (t, c) in enumerate(steps):
    x = x0 + i * (bw + gap)
    box(s, x, 1.9, bw, 1.3, t, fill=c, size=17)
    if i < len(steps) - 1:
        arrow(s, x + bw + 0.05, 2.4, gap - 0.1, 0.3)
bullets(s, [
    "Filtre du bruit : contour de longueur > 20 px et étendue > 30 px (flèches d'axe, restes de lettres)",
    "Piège : cv2.contourArea vaut 0 pour une ligne d'1 pixel \u2192 on utilise cv2.arcLength",
    "Les bords anti-aliasés des courbes colorées auraient formé de fausses courbes noires \u2192 masque dilaté",
    "Les courbes colorées et la courbe noire coexistent dans la même image",
], 0.7, 3.7, 12, 3.3, size=21)

# ------------------------------------------------------------------ 7. legendes + export
s = new_slide("Étape 4 \u2013 Légendes et export", "Nommer les courbes, convertir, sauvegarder", notes=(
    "Idée inspirée de GraphMaster : pour chaque texte qui n'est pas un nombre, on cherche une pastille de couleur "
    "à sa gauche. Un vrai carré ou trait de légende est plein ; les chiffres et lettres sont rejetés par leur forme. "
    "Ensuite appariement un-pour-un avec les courbes par distance de couleur. Limite : si la légende est écrite en "
    "texte (Dots / Lines) il n'y a pas de pastille, donc nom générique. Note : le dépôt GraphMaster public ne contient "
    "que le README, pas le code : on s'est inspiré de la méthode décrite, pas du code."))
bullets(s, [
    "Légende : texte non numérique + pastille pleine (carré ou trait) à sa gauche",
    ("chiffres et lettres rejetés par leur forme (remplissage, hauteur)", 1),
    "Appariement un-pour-un courbe \u2194 légende par distance de couleur RGB",
    "Sans légende exploitable : nom générique (Courbe_1, Courbe_2, …)",
    "Pixels \u2192 valeurs : moyenne des y pour chaque x (trait épais \u2192 1 point), puis conversion par l'échelle",
    "Sortie : un .xlsx (colonnes X/Y par courbe) + une image de vérification côte à côte",
], 0.7, 1.55, 12, 5.4, size=22)

# ------------------------------------------------------------------ 8-9. resultats
def result_slide(title, sub, img, points, notes):
    s = new_slide(title, sub, notes=notes)
    picture(s, os.path.join(DATA, img), 0.4, 1.5, 8.7, 5.6)
    bullets(s, points, 9.3, 1.9, 3.8, 5, size=18)
    return s


result_slide("Résultat \u2013 trois courbes dans un graphe", "test_graph.png (couleurs + noir)",
             "test_graph_verification.png", [
                 "Courbes bleue, rouge et noire séparées",
                 "Échelle X 0 \u2192 10 et Y 0 \u2192 1 lues par l'OCR",
                 "Texte « 750 °C » et légende ignorés",
                 "Légende en texte (Dots / Lines) : pas de pastille \u2192 noms génériques",
             ], "Cas le plus complet. La calibration X était fausse (5 \u2192 10) avant l'ajustement robuste ; "
                "elle est maintenant correcte. Le point noir : la courbe noire mélange points et trait.")

result_slide("Résultat \u2013 nuage de points", "Data_points.png",
             "Data_points_verification.png", [
                 "Une seule série, bleue, détectée",
                 "Échelle X 10 \u2192 60 et Y 0 \u2192 1 correctes",
                 "Les points sont reliés par une ligne : on récupère la tendance, pas les points individuels",
                 "Effet d'escalier : une colonne de pixels donne un seul y",
             ], "Les points isolés deviennent une courbe en escalier. Piste d'amélioration : détecter les centres "
                "de marqueurs (composantes connexes) pour restituer les points réels.")

# ------------------------------------------------------------------ 10. cas limite
s = result_slide("Cas limite \u2013 axes en fractions de π", "image082.png : sin θ, graduations 0, π/2, π, 3π/2, 2π",
                 "image082_verification.png", [
                     "L'OCR lit « 2 » et « 3 » dans « 2π » et « 3π/2 » : pas de vraies graduations numériques",
                     "Avant : une échelle fausse était acceptée",
                     "Maintenant : calibration refusée \u2192 valeurs en pixels, signalé",
                     "La forme de la courbe reste correcte",
                 ], "Exemple de limite honnête. On préfère un résultat en pixels, clairement signalé, plutôt qu'un "
                    "résultat qui a l'air juste mais dont l'échelle est fausse. Une amélioration possible : reconnaître π.")

# ------------------------------------------------------------------ 10b. Code - RANSAC
s = new_slide("Sous le capot : Calibration robuste", "Extrait de main.py (fit_axis)", notes="Explication de la logique RANSAC")
bullets(s, [
    "Rejet des valeurs OCR aberrantes par consensus :",
    "On teste toutes les paires de points pour créer une droite",
    "On compte combien de points valident cette droite",
    "On garde la droite qui a le plus de consensus"
], 0.5, 1.5, 12, 1.5, size=20)
code_ransac = '''def fit_axis(points, key):
    # ...
    best_score, best_inl = None, None
    for i in range(len(pos)):
        for j in range(i + 1, len(pos)):
            if abs(pos[i] - pos[j]) < 5 or val[i] == val[j]: continue
            
            # Droite passant par les points i et j
            a = (val[j] - val[i]) / (pos[j] - pos[i])
            b = val[i] - a * pos[i]
            
            # Comptage des inliers (tolérance de 4%)
            inl = np.abs(a * pos + b - val) <= tol
            score = (int(inl.sum()), float(np.ptp(pos[inl])))
            
            if best_score is None or score > best_score:
                best_score, best_inl = score, inl
                
    if int(best_inl.sum()) < 3: return None, 0 # Refus si pas fiable
    # ...'''
code_box(s, 1.0, 3.2, 11.0, 4.0, code_ransac, size=14)

# ------------------------------------------------------------------ 10c. Code - Clustering
s = new_slide("Sous le capot : Séparation des couleurs", "Extrait de main.py (detect_curves)", notes="Explication de la logique MeanShift")
bullets(s, [
    "Clustering sur la teinte (HSV) :",
    "La teinte est circulaire (rouge à 0° et 179°)",
    "Projection cosinus/sinus pour éviter la cassure à 180°",
    "MeanShift découvre tout seul le nombre de courbes"
], 0.5, 1.5, 12, 1.5, size=20)
code_color = '''def detect_curves(img, text_boxes):
    hsv = cv2.cvtColor(img, cv2.COLOR_BGR2HSV)
    mask = (hsv[:, :, 1] > 50) & (hsv[:, :, 2] > 50)  # Pixels colorés
    
    # Teinte = angle sur un cercle (0-180 dans OpenCV)
    # Projection (cos, sin) pour que rouge (0) et rouge (179) soient voisins
    ang = hsv[:, :, 0][mask] * (np.pi / 90.0)
    feats = np.column_stack((np.cos(ang) * 100, np.sin(ang) * 100))
    
    # Clustering automatique avec MeanShift
    sample = rng.choice(len(feats), min(3000, len(feats)), replace=False)
    labels = MeanShift(bandwidth=30, bin_seeding=True).fit(feats[sample]).predict(feats)
    
    for lab in np.unique(labels):
        sel = labels == lab
        if sel.sum() >= MIN_CLUSTER_PIXELS:
            curves.append({"coords": coords[sel], "color": pix_rgb[sel].mean(axis=0)})'''
code_box(s, 1.0, 3.2, 11.0, 4.0, code_color, size=14)

# ------------------------------------------------------------------ 10d. Code - Black Curve
s = new_slide("Sous le capot : Extraction de la courbe noire", "Extrait de main.py (extract_black_curves)", notes="Explication de la morphologie mathématique")
bullets(s, [
    "Extraction par soustraction :",
    "On efface le texte à l'aide des boîtes OCR",
    "On efface les axes avec la morphologie mathématique",
    "On efface les courbes colorées avec un masque dilaté"
], 0.5, 1.5, 12, 1.5, size=20)
code_black = '''def extract_black_curves(image, text_boxes, mask_colored):
    gray = cv2.cvtColor(image, cv2.COLOR_BGR2GRAY)
    _, thresh = cv2.threshold(gray, 150, 255, cv2.THRESH_BINARY_INV)

    # 1. Gommer le texte
    for x0, y0, x1, y1 in text_boxes:
        cv2.rectangle(thresh, (int(x0)-5, int(y0)-5), (int(x1)+5, int(y1)+5), 0, -1)

    # 2. Gommer lignes droites (axes, grille) via ouverture morphologique
    hl = cv2.morphologyEx(thresh, cv2.MORPH_OPEN, cv2.getStructuringElement(cv2.MORPH_RECT, (40, 1)))
    vl = cv2.morphologyEx(thresh, cv2.MORPH_OPEN, cv2.getStructuringElement(cv2.MORPH_RECT, (1, 40)))
    curve_only = cv2.subtract(thresh, cv2.dilate(cv2.add(hl, vl), np.ones((3,3))))

    # 3. Gommer courbes colorées (avec dilatation pour les bords anti-aliasés)
    colored = cv2.dilate(mask_colored.astype(np.uint8)*255, np.ones((3,3)), iterations=3)
    curve_only = cv2.subtract(curve_only, colored)'''
code_box(s, 1.0, 3.2, 11.0, 4.0, code_black, size=14)

# ------------------------------------------------------------------ 11. difficultes
s = new_slide("Difficultés rencontrées et solutions", "Ce que le développement nous a appris", notes=(
    "Slide à raconter comme un retour d'expérience. Le plus formateur : les erreurs silencieuses, par exemple une "
    "légende qui était en réalité un titre d'axe, ou un mélange RGB/BGR dans la comparaison des couleurs."))
rows = [
    ("Chiffres d'axes illisibles par l'OCR", "Agrandissement ×4 + binarisation (méthode du dépôt GTE.py)"),
    ("Courbe noire fusionnée avec texte/grille", "Soustraction : boîtes OCR, morphologie, masque coloré"),
    ("Lignes fines supprimées par erreur", "contourArea = 0 pour 1 pixel \u2192 arcLength"),
    ("Bords anti-aliasés = fausses courbes", "Dilatation des masques avant soustraction"),
    ("Échelle OCR fausse (chiffre mal lu)", "Ajustement robuste + refus si peu fiable"),
    ("Titre d'axe pris pour une légende", "Légende = texte + vraie pastille (forme testée)"),
    ("Comparaison de couleurs RGB vs BGR", "Tout converti en RGB avant comparaison"),
]
y = 1.5
for a, b in rows:
    box(s, 0.5, y, 5.6, 0.62, a, fill=LIGHT, size=15, color=NAVY, bold=False)
    arrow(s, 6.2, y + 0.13, 0.4, 0.36, TEAL)
    box(s, 6.7, y + 0.0, 6.1, 0.62, b, fill=TEAL, size=15, color=WHITE, bold=False)
    y += 0.76

# ------------------------------------------------------------------ 12. optimisation
s = new_slide("Optimisation du code", "Réécriture de main.py", notes=(
    "Honnêteté : une partie du code était dupliquée à cause d'un mauvais assemblage de fichiers ; elle a été supprimée. "
    "Les gains de vitesse viennent surtout de la vectorisation et du plafonnement de l'agrandissement. "
    "Le temps total reste dominé par l'OCR, qui tourne sur CPU."))
bullets(s, [
    "Code mort et dupliqué supprimé (fonction de légendes présente deux fois), scripts de test archivés",
    "Moyenne des y par colonne vectorisée (np.unique + bincount) au lieu d'une boucle Python par colonne",
    "Concaténation unique des coordonnées au lieu d'étendre des listes de tableaux",
    "OCR : agrandissement plafonné à 2560 px, chargé à la demande, 2e lecture seulement si nécessaire",
    "Fonctions courtes à responsabilité unique (OCR, calibration, légendes, courbes, export)",
    "Un lot ne s'arrête plus si une image échoue ; résultats reproductibles (graine fixée)",
    "requirements.txt complété (matplotlib, easyocr, openpyxl manquaient)",
], 0.7, 1.55, 12, 5.4, size=21)

# ------------------------------------------------------------------ 13. limites / perspectives
s = new_slide("Limites et perspectives", "Ce qui reste à faire", notes=(
    "Les trois premières perspectives sont déjà identifiées : échelles logarithmiques, courbes qui ne sont pas des "
    "fonctions de x (boucles, spirales) avec suivi de trajectoire, et deux axes Y. "
    "Aucun taux de précision chiffré n'est présenté : la validation actuelle est visuelle."))
box(s, 0.5, 1.6, 6.0, 0.6, "Limites actuelles", fill=ORANGE, size=20)
bullets(s, [
    "Axes linéaires uniquement",
    "Une seule valeur y par x : pas de boucles ni de spirales",
    "Un seul axe Y",
    "Légendes textuelles (Dots / Lines) non exploitées",
    "Graduations non numériques (π, dates) non lues",
    "Validation visuelle : pas encore de taux de précision mesuré",
], 0.6, 2.4, 6.0, 4.6, size=19)
box(s, 6.9, 1.6, 6.0, 0.6, "Perspectives", fill=TEAL, size=20)
bullets(s, [
    "Échelles logarithmiques : détecter l'espacement des graduations",
    "Suivi de trajectoire pour les courbes non bijectives",
    "Deux axes Y (gauche / droite)",
    "Marqueurs : retrouver les points réels au lieu d'une ligne",
    "Mesurer la précision sur des graphes de référence connus",
], 7.0, 2.4, 6.0, 4.6, size=19)

# ------------------------------------------------------------------ 14. conclusion
s = new_slide("Conclusion", notes=(
    "Message final : la chaîne complète fonctionne sans aucune saisie, elle signale ses limites plutôt que de les "
    "cacher, et elle est prête à être étendue."))
bullets(s, [
    "Pipeline complet et automatique : image \u2192 Excel + image de contrôle",
    "Plusieurs courbes (colorées et noires) séparées et nommées quand la légende le permet",
    "Calibration robuste : refuse plutôt que d'inventer une échelle",
    "Inspirations : GTE.py (prétraitement OCR) et la méthode GraphMaster (couleurs + légendes)",
    "Prochaines étapes : axes logarithmiques, courbes non bijectives, double axe Y, mesure de précision",
], 0.7, 1.7, 12, 4.5, size=24)
text_box(s, 0.7, 6.2, 12, 0.6, "Merci \u2013 questions ?", size=28, color=TEAL, bold=True)

prs.save(OUT)
print("Sauvegarde :", OUT)
